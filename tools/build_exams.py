"""Build the practice-exam downloads from the reviewed question banks.

    python tools/build_exams.py            # rebuilds exams/core1 and exams/core2

Source of truth: exams/src/core1.json and exams/src/core2.json
(each question: domain, objective, q, correct[], wrong[], exp).

Answer positions are shuffled at build time with a fixed seed and tracked by option TEXT, so the
paper and the key always agree and the correct answer is not biased toward any one letter.
Needs: pip install python-docx
"""

import json
import random
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
LETTERS = ["A", "B", "C", "D"]
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x60, 0x60, 0x60)
CORRECT = RGBColor(0x1E, 0x6B, 0x3A)
CHOOSE_RE = re.compile(r"\s*\((?:choose|select)\s+(?:two|2)\.?\)\s*$", re.I)

EXAMS = {
    "core1": {
        "title": "CompTIA A+ Core 1", "code": "220-1201", "name": "Practice Exam 1", "seed": 20261004,
        "passing": "675 out of 900",
        "weights": {"Mobile Devices": 13, "Networking": 23, "Hardware": 25,
                    "Virtualization and Cloud Computing": 11, "Hardware and Network Troubleshooting": 28},
    },
    "core2": {
        "title": "CompTIA A+ Core 2", "code": "220-1202", "name": "Practice Exam 1", "seed": 20261004,
        "passing": "700 out of 900",
        "weights": {"Operating Systems": 28, "Security": 28, "Software Troubleshooting": 23,
                    "Operational Procedures": 21},
    },
}


def shuffle(questions, seed):
    rng = random.Random(seed)
    out = []
    for i, q in enumerate(questions, start=1):
        options = list(q["correct"]) + list(q["wrong"])
        if len(options) != 4 or len(set(options)) != 4:
            raise ValueError(f"Q{i}: need 4 distinct options, got {options}")
        rng.shuffle(options)
        lettered = dict(zip(LETTERS, options))
        letters = sorted(L for L, t in lettered.items() if t in q["correct"])
        if len(letters) != len(q["correct"]):
            raise ValueError(f"Q{i}: correct answers lost in shuffle")
        out.append({"number": i, "domain": q["domain"], "objective": q.get("objective", ""),
                    "question": CHOOSE_RE.sub("", q["q"]).strip(), "options": lettered,
                    "letters": letters, "multi": len(q["correct"]) > 1, "exp": q["exp"]})
    return out


def base_doc():
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Pt(52)
        s.left_margin = s.right_margin = Pt(60)
    return d


def centered(doc, text, size, color=None, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    if color is not None:
        r.font.color.rgb = color


def heading(doc, text):
    r = doc.add_paragraph().add_run(text)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(13), ACCENT


def guide_lines(n, cfg):
    hi, mid = round(n * 0.80), round(n * 0.70)
    return [f"{hi} or more out of {n}: comfortably ready, if your practice-exam scores are consistently this high.",
            f"{mid} to {hi - 1}: borderline. Drill the domain with the most misses, then retake a different paper.",
            f"Below {mid}: not ready yet. Go back to the parts that cover your misses, using the objective numbers."]


def build_docx(items, cfg, counts, qpath, apath):
    n, multi = len(items), sum(i["multi"] for i in items)
    d = base_doc()
    centered(d, f"{cfg['title']} ({cfg['code']})", 24, ACCENT, True)
    centered(d, f"{cfg['name']}: Question Paper", 15, MUTED)
    centered(d, f"{n} questions  ·  90 minutes  ·  no notes, no lookups", 10, MUTED)
    d.add_paragraph()
    heading(d, "Instructions")
    for line in [
        f"Timed at 90 minutes, the real exam's limit. The real exam has at most 90 questions; this paper has {n}.",
        f"The real exam pass mark is {cfg['passing']}. CompTIA uses scaled scoring, so percentages here are only a rough guide.",
        f"{multi} questions ask for TWO answers and are marked. Both must be right to score.",
        "Answer every question; a blank scores nothing.",
        "The real exam also has performance-based questions (hands-on tasks). This paper is multiple choice only: use the labs in this study kit for hands-on practice.",
        "Answers and explanations are in the separate answer-key file.",
    ]:
        d.add_paragraph(line, style="List Bullet").paragraph_format.space_after = Pt(3)
    d.add_paragraph()
    heading(d, "Domain coverage (follows the published exam weights)")
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for c, txt in zip(t.rows[0].cells, ["Domain", "Questions", "Exam weight"]):
        c.text = txt
    for dom, w in cfg["weights"].items():
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text = dom, str(counts[dom]), f"{w}%"
    d.add_page_break()
    for it in items:
        p = d.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(4)
        num = p.add_run(f"{it['number']}.  ")
        num.bold, num.font.color.rgb = True, ACCENT
        p.add_run(it["question"])
        if it["multi"]:
            p.add_run("  (Choose TWO.)").bold = True
        for L in LETTERS:
            o = d.add_paragraph()
            o.paragraph_format.left_indent, o.paragraph_format.space_after = Pt(24), Pt(2)
            o.add_run(f"{L}.  ").bold = True
            o.add_run(it["options"][L])
    d.save(qpath)

    a = base_doc()
    centered(a, f"{cfg['title']} ({cfg['code']})", 24, ACCENT, True)
    centered(a, f"{cfg['name']}: Answer Key and Explanations", 15, MUTED)
    centered(a, "Keep this separate from the question paper while you sit the exam", 10, MUTED)
    a.add_paragraph()
    heading(a, "Quick key")
    tbl = a.add_table(rows=0, cols=10)
    tbl.style = "Light Grid Accent 1"
    row = None
    for idx, it in enumerate(items):
        if idx % 10 == 0:
            row = tbl.add_row().cells
        row[idx % 10].text = f"{it['number']}. {''.join(it['letters'])}"
    a.add_paragraph()
    heading(a, "Score yourself")
    for line in guide_lines(n, cfg):
        a.add_paragraph(line, style="List Bullet")
    a.add_paragraph("Track the domain of each miss. A cluster in one domain tells you far more than the total.",
                    style="List Bullet")
    a.add_page_break()
    heading(a, "Explanations")
    for it in items:
        p = a.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(11), Pt(2)
        num = p.add_run(f"{it['number']}.  ")
        num.bold, num.font.color.rgb = True, ACCENT
        ans = p.add_run("Answer: " + ", ".join(it["letters"]))
        ans.bold, ans.font.color.rgb = True, CORRECT
        tag = f"   [{it['domain']}" + (f", objective {it['objective']}" if it["objective"] else "") + "]"
        t2 = p.add_run(tag)
        t2.font.size, t2.font.color.rgb = Pt(9), MUTED
        for L in it["letters"]:
            o = a.add_paragraph()
            o.paragraph_format.left_indent, o.paragraph_format.space_after = Pt(24), Pt(2)
            r = o.add_run(f"{L}.  {it['options'][L]}")
            r.font.color.rgb = CORRECT
        e = a.add_paragraph()
        e.paragraph_format.left_indent = Pt(24)
        e.add_run(it["exp"]).font.size = Pt(10)
    a.save(apath)


def build_md(items, cfg, counts, qpath, apath):
    n = len(items)
    q = [f"# {cfg['title']} ({cfg['code']}): {cfg['name']}", "",
         f"{n} questions, 90 minutes. Pass mark on the real exam: {cfg['passing']} (scaled). Answers are in a separate file.", ""]
    for it in items:
        q += [f"**{it['number']}.** {it['question']}" + ("  *(Choose TWO.)*" if it["multi"] else ""), ""]
        q += [f"- {L}. {it['options'][L]}" for L in LETTERS] + [""]
    Path(qpath).write_text("\n".join(q), encoding="utf-8")
    a = [f"# {cfg['title']} ({cfg['code']}): {cfg['name']}, answer key", "", "## Quick key", ""]
    a += ["| " + " | ".join(f"{i['number']}: {''.join(i['letters'])}" for i in items[s:s + 10]) + " |"
          for s in range(0, n, 10)]
    a += ["", "## Score yourself", ""] + [f"- {x}" for x in guide_lines(n, cfg)] + ["", "## Explanations", ""]
    for it in items:
        tag = it["domain"] + (f", objective {it['objective']}" if it["objective"] else "")
        a += [f"**{it['number']}. Answer: {', '.join(it['letters'])}**  *[{tag}]*", ""]
        a += [f"- {L}. {it['options'][L]}" for L in it["letters"]] + ["", it["exp"], ""]
    Path(apath).write_text("\n".join(a), encoding="utf-8")


def main():
    for key, cfg in EXAMS.items():
        src = json.loads((ROOT / "exams" / "src" / f"{key}.json").read_text(encoding="utf-8"))
        items = shuffle(src, cfg["seed"])
        counts = Counter(i["domain"] for i in items)
        unknown = set(counts) - set(cfg["weights"])
        if unknown:
            sys.exit(f"{key}: unknown domains {unknown}")
        out = ROOT / "exams" / key
        out.mkdir(parents=True, exist_ok=True)
        build_docx(items, cfg, counts, out / "practice-exam-1-questions.docx", out / "practice-exam-1-answers.docx")
        build_md(items, cfg, counts, out / "practice-exam-1-questions.md", out / "practice-exam-1-answers.md")
        pos = Counter(L for i in items for L in i["letters"])
        print(f"{key}: {len(items)} questions, {sum(i['multi'] for i in items)} choose-two, "
              f"key letters {dict(sorted(pos.items()))}, domains {dict(counts)}")


if __name__ == "__main__":
    main()
